import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_thesis_form():
    doc = docx.Document()
    
    # Page setup - standard 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles & Fonts
    primary_color = RGBColor(27, 54, 93)     # Deep Navy
    text_color = RGBColor(40, 40, 40)        # Dark Charcoal
    light_bg = "F4F6F9"

    # Set base normal style font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = text_color

    # Title Banner
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("THESIS REGISTRATION FORM")
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = primary_color
    title_p.paragraph_format.space_after = Pt(12)

    # Instructions Box
    inst_p = doc.add_paragraph()
    inst_p.paragraph_format.space_after = Pt(14)
    inst_run = inst_p.add_run(
        "INSTRUCTIONS:\n"
        "• Student: Please fill your Name and ID on the form and discuss with your Thesis Advisor to fill the other contents. Afterward, sign the form to verify your agreement and return the form to the Undergraduate Academic Assistant of the Department.\n"
        "• Thesis Supervisor: Please authorize the Thesis Registration Form for this student by signing below."
    )
    inst_run.font.italic = True
    inst_run.font.size = Pt(9.5)
    inst_run.font.color.rgb = RGBColor(80, 80, 80)

    # Helper function to set table cell background color
    def set_cell_background(cell, fill_hex):
        shading_xml = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading_xml)

    # Helper function to set cell padding
    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    # Table 1: Student & Project Metadata Information
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Widths: Col 0 = 2.0 inches, Col 1 = 4.5 inches
    col_widths = [Inches(2.0), Inches(4.5)]
    
    metadata = [
        ("Student's Name:", "[Fill Student Name]"),
        ("Student ID:", "[Fill Student ID]"),
        ("Email:", "[Fill Email Address]"),
        ("Phone:", "[Fill Phone Number]"),
        ("Major:", "Computer Science / Software Engineering"),
    ]

    for row_idx, (label, val) in enumerate(metadata):
        row = table.rows[row_idx]
        
        # Label cell
        cell_lbl = row.cells[0]
        cell_lbl.width = col_widths[0]
        set_cell_background(cell_lbl, light_bg)
        set_cell_margins(cell_lbl, top=100, bottom=100, left=150, right=150)
        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_after = Pt(0)
        run_lbl = p_lbl.add_run(label)
        run_lbl.font.bold = True
        run_lbl.font.color.rgb = primary_color

        # Value cell
        cell_val = row.cells[1]
        cell_val.width = col_widths[1]
        set_cell_margins(cell_val, top=100, bottom=100, left=150, right=150)
        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_after = Pt(0)
        run_val = p_val.add_run(val)
        if "[" in val:
            run_val.font.italic = True
            run_val.font.color.rgb = RGBColor(120, 120, 120)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section Header Helper
    def add_section_header(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = primary_color

    # Thesis Title
    add_section_header("THESIS TITLE")
    title_box = doc.add_table(rows=1, cols=1)
    title_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_t = title_box.rows[0].cells[0]
    cell_t.width = Inches(6.5)
    set_cell_background(cell_t, light_bg)
    set_cell_margins(cell_t, top=140, bottom=140, left=180, right=180)
    p_t = cell_t.paragraphs[0]
    p_t.paragraph_format.space_after = Pt(0)
    run_t = p_t.add_run("CogniTrace: An Intelligent AI Tutoring System for Programming Pedagogy via State-Grounded LLMs, Animated Memory Graphs, and Time-Travel Execution Forking")
    run_t.font.bold = True
    run_t.font.size = Pt(11.5)
    run_t.font.color.rgb = primary_color

    # Thesis Goals and Objectives
    add_section_header("THESIS GOALS AND OBJECTIVES")
    
    goals_p = doc.add_paragraph()
    goals_p.paragraph_format.space_after = Pt(4)
    goals_p.add_run("The primary objective of this thesis is to design, implement, and evaluate CogniTrace—an Intelligent AI Tutoring System for computer science learners that transforms passive code execution visualization into an active, interactive, and personalized learning experience. Specific goals include:").font.size = Pt(10.5)

    objectives = [
        ("Animated Memory Heap Graphs", "Render interactive, animated node-link visualizations of Python object memory addresses (id()), reference aliasing, and pointer mutations in real time to resolve fundamental mental model bottlenecks in programming education."),
        ("Time Scrubbing (Execution Scrubbing)", "Provide a continuous, video-style timeline scrubber enabling students to smoothly navigate back and forth across execution history, inspecting micro-level variable state changes."),
        ("State-Forking (Parallel Timeline Branching)", "Enable students to scrub back to any execution step, modify memory variable values mid-trace, and execute parallel 'What-If' execution timelines to visually compare divergent state outcomes."),
        ("State-Grounded AI Guidance", "Integrate a dual-LLM reasoning engine (Ollama/OpenAI) that generates grounded, line-level explanations and 3-tier Socratic hint ladders based on live runtime telemetry (memory IDs, AST nodes, bytecode) rather than static code descriptions."),
        ("Active Execution Prediction Engine", "Develop a real-time execution interception pipeline using Python's sys.settrace() that pauses execution at critical control-flow points, prompting learners to predict variable states before revealing answers.")
    ]

    for title, desc in objectives:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        bp.paragraph_format.space_before = Pt(0)
        r_bold = bp.add_run(f"{title}: ")
        r_bold.font.bold = True
        r_bold.font.color.rgb = primary_color
        r_desc = bp.add_run(desc)
        r_desc.font.size = Pt(10)

    # Requirements
    add_section_header("REQUIREMENTS")

    req_intro = doc.add_paragraph()
    req_intro.paragraph_format.space_after = Pt(4)
    req_intro.add_run("The thesis project will fulfill the following technical, functional, and non-functional requirements:").font.size = Pt(10.5)

    requirements = [
        ("Functional Requirements", [
            "Animated Memory Heap Graph Component: Front-end canvas (React Flow / D3.js) rendering animated object nodes, reference pointer links, and garbage collection states driven by Python id() addresses.",
            "Time Scrubbing & State-Forking UI: Slider UI supporting smooth step scrubbing, mid-trace variable state mutation, and dual-trace side-by-side execution diffing.",
            "State-Grounded AI Explanation Engine: LLM prompt pipeline anchored directly in sys.settrace() runtime frame variables and bytecode state.",
            "Execution Telemetry Sandbox: Pure computational sandbox blocking unapproved filesystem, network, and subprocess calls during Python trace execution.",
            "Socratic Hint Ladder: 3-tier progressive guidance mode (Nudge -> Clue -> Complete Explanation) to mitigate AI-code over-reliance."
        ]),
        ("Non-Functional & Technical Requirements", [
            "Backend Architecture: High-performance FastAPI server with Server-Sent Events (SSE) for streaming line-by-line trace telemetry and LLM explanations.",
            "Frontend Architecture: React + Vite + TypeScript + Tailwind CSS design system with smooth micro-animations and zero ad-hoc styling.",
            "LLM Integration: Hybrid routing via local Ollama Cloud / OpenAI API with fallback error handling.",
            "Execution Performance: Trace overhead latency under 150ms per line step and instant visual graph re-renders."
        ])
    ]

    for req_cat, req_list in requirements:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after = Pt(2)
        r_cat = cp.add_run(req_cat)
        r_cat.font.bold = True
        r_cat.font.size = Pt(11)
        r_cat.font.color.rgb = primary_color

        for item in req_list:
            bp = doc.add_paragraph(style='List Bullet')
            bp.paragraph_format.space_after = Pt(2)
            parts = item.split(":", 1)
            if len(parts) == 2:
                r_b = bp.add_run(f"{parts[0]}:")
                r_b.font.bold = True
                r_t = bp.add_run(parts[1])
            else:
                r_t = bp.add_run(item)
            bp.runs[0].font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Signatures Section
    add_section_header("AUTHORIZATION AND SIGNATURES")

    sig_table = doc.add_table(rows=2, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False

    sig_col_widths = [Inches(3.25), Inches(3.25)]

    sig_data = [
        ("Student Signature:", "Thesis Supervisor Signature:"),
        ("Date: ________________________", "Date: ________________________")
    ]

    for row_idx, (col1, col2) in enumerate(sig_data):
        row = sig_table.rows[row_idx]
        
        c1 = row.cells[0]
        c1.width = sig_col_widths[0]
        set_cell_margins(c1, top=100, bottom=100, left=100, right=100)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(6)
        r1 = p1.add_run(col1)
        r1.font.size = Pt(10.5)
        if row_idx == 0:
            r1.font.bold = True
            r1.font.color.rgb = primary_color
            p1.add_run("\n\n_____________________________________\n")

        c2 = row.cells[1]
        c2.width = sig_col_widths[1]
        set_cell_margins(c2, top=100, bottom=100, left=100, right=100)
        p2 = c2.paragraphs[0]
        p2.paragraph_format.space_after = Pt(6)
        r2 = p2.add_run(col2)
        r2.font.size = Pt(10.5)
        if row_idx == 0:
            r2.font.bold = True
            r2.font.color.rgb = primary_color
            p2.add_run("\n\n_____________________________________\n")

    # Save to workspace root
    output_filename = "Thesis_Registration_Form.docx"
    doc.save(output_filename)
    print(f"Successfully generated {output_filename}")

if __name__ == "__main__":
    create_thesis_form()
